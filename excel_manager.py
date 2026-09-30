from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

from database import (
    get_dashboard_summary,
    list_employees,
    list_equipment,
    list_transport_vehicles,
    list_water_trucks,
    get_employee_records,
    get_equipment_records,
    get_transport_records,
    get_water_records,
    get_employee,
    get_equipment,
    get_transport_vehicle,
    get_water_truck,
)


OUTPUT_DIR = Path(__file__).resolve().parent / "exports"
OUTPUT_DIR.mkdir(exist_ok=True)


def safe_sheet_name(name: str, max_len=31):
    clean = "".join(ch for ch in str(name) if ch.isalnum() or ch in " _-")
    clean = clean[:max_len]
    return clean or "Sheet"


def style_header(ws, row=1, col_start=1, col_end=None):
    if col_end is None:
        col_end = ws.max_column
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for col in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(
            left=Side(style="thin", color="D9E2F3"),
            right=Side(style="thin", color="D9E2F3"),
            top=Side(style="thin", color="D9E2F3"),
            bottom=Side(style="thin", color="D9E2F3"),
        )


def apply_table_style(ws):
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            if cell.row > 1 and cell.column > 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")


def add_summary_sheet(wb):
    ws = wb.active
    ws.title = "الرئيسية"
    ws.freeze_panes = "A2"
    ws["A1"] = "لوحة التحكم"
    ws["A1"].font = Font(size=16, bold=True, color="1F1F1F")

    summary = get_dashboard_summary()
    stats = [
        ("عدد العمال", summary["employees_count"]),
        ("عدد المعدات", summary["equipment_count"]),
        ("عدد عربيات النقل", summary["transport_count"]),
        ("عدد عربيات المياه", summary["water_count"]),
        ("إجمالي مستحقات العمال", summary["employee_net_total"]),
        ("إجمالي تشغيل المعدات", summary["equipment_total"]),
        ("إجمالي عربيات النقل", summary["transport_total"]),
        ("إجمالي عربيات المياه", summary["water_total"]),
    ]

    ws["A3"] = "العنوان"
    ws["B3"] = "القيمة"
    style_header(ws, 3, 1, 2)
    for i, (label, value) in enumerate(stats, start=4):
        ws[f"A{i}"] = label
        ws[f"B{i}"] = value
    for col in ["A", "B"]:
        ws.column_dimensions[col].width = 30
    ws.auto_filter.ref = "A3:B" + str(len(stats) + 3)


def export_employee_detail(ws, employee):
    ws["A1"] = "الكود"
    ws["B1"] = employee["code"]
    ws["A2"] = "الاسم"
    ws["B2"] = employee["name"]
    ws["A3"] = "الوظيفة"
    ws["B3"] = employee["job_title"]
    ws["A4"] = "الهاتف"
    ws["B4"] = employee["phone"]
    ws["A5"] = "المشروع"
    ws["B5"] = employee["project"]
    ws["A6"] = "سعر اليومية"
    ws["B6"] = employee["daily_rate"]

    ws["A8"] = "التاريخ"
    ws["B8"] = "المشروع"
    ws["C8"] = "الحالة"
    ws["D8"] = "اليومية"
    ws["E8"] = "إضافي"
    ws["F8"] = "خصم"
    ws["G8"] = "صافي المستحق"
    ws["H8"] = "ملاحظات"
    style_header(ws, 8, 1, 8)
    rows = get_employee_records(employee["id"])
    for idx, row in enumerate(rows, start=9):
        ws[f"A{idx}"] = row["record_date"]
        ws[f"B{idx}"] = row["project"]
        ws[f"C{idx}"] = row["status"]
        ws[f"D{idx}"] = row["daily_amount"]
        ws[f"E{idx}"] = row["extra"]
        ws[f"F{idx}"] = row["discount"]
        ws[f"G{idx}"] = row["net_amount"]
        ws[f"H{idx}"] = row["notes"]

    ws["A{0}".format(len(rows)+11)] = "إجمالي اليوميات"
    ws["B{0}".format(len(rows)+11)] = sum(float(r["daily_amount"]) for r in rows)
    ws["A{0}".format(len(rows)+12)] = "إجمالي الإضافي"
    ws["B{0}".format(len(rows)+12)] = sum(float(r["extra"]) for r in rows)
    ws["A{0}".format(len(rows)+13)] = "إجمالي الخصومات"
    ws["B{0}".format(len(rows)+13)] = sum(float(r["discount"]) for r in rows)
    ws["A{0}".format(len(rows)+14)] = "صافي المستحق"
    ws["B{0}".format(len(rows)+14)] = sum(float(r["net_amount"]) for r in rows)

    for cell in ws["A1":"H" + str(ws.max_row)]:
        for c in cell:
            c.border = Border(
                left=Side(style="thin", color="D9E2F3"),
                right=Side(style="thin", color="D9E2F3"),
                top=Side(style="thin", color="D9E2F3"),
                bottom=Side(style="thin", color="D9E2F3"),
            )
    ws.freeze_panes = "A9"


def export_equipment_detail(ws, equipment):
    ws["A1"] = "كود المعدة"
    ws["B1"] = equipment["code"]
    ws["A2"] = "اسم المعدة"
    ws["B2"] = equipment["name"]
    ws["A3"] = "النوع"
    ws["B3"] = equipment["equipment_type"]
    ws["A4"] = "رقم المعدة"
    ws["B4"] = equipment["equipment_number"]
    ws["A5"] = "المشغل"
    ws["B5"] = equipment["operator_name"]
    ws["A6"] = "المشروع"
    ws["B6"] = equipment["project"]
    ws["A7"] = "سعر الساعة"
    ws["B7"] = equipment["hourly_rate"]

    ws["A9"] = "التاريخ"
    ws["B9"] = "المشروع"
    ws["C9"] = "بداية التشغيل"
    ws["D9"] = "نهاية التشغيل"
    ws["E9"] = "عدد الساعات"
    ws["F9"] = "سعر الساعة"
    ws["G9"] = "إجمالي المستحق"
    ws["H9"] = "ملاحظات"
    style_header(ws, 9, 1, 8)
    rows = get_equipment_records(equipment["id"])
    for idx, row in enumerate(rows, start=10):
        ws[f"A{idx}"] = row["record_date"]
        ws[f"B{idx}"] = row["project"]
        ws[f"C{idx}"] = row["start_time"]
        ws[f"D{idx}"] = row["end_time"]
        ws[f"E{idx}"] = row["hours"]
        ws[f"F{idx}"] = row["hourly_rate"]
        ws[f"G{idx}"] = row["total_amount"]
        ws[f"H{idx}"] = row["notes"]
    ws.freeze_panes = "A10"


def export_transport_detail(ws, vehicle):
    ws["A1"] = "الكود"
    ws["B1"] = vehicle["code"]
    ws["A2"] = "رقم العربية"
    ws["B2"] = vehicle["vehicle_number"]
    ws["A3"] = "نوع العربية"
    ws["B3"] = vehicle["vehicle_type"]
    ws["A4"] = "لوحة العربية"
    ws["B4"] = vehicle["plate_number"]
    ws["A5"] = "السائق"
    ws["B5"] = vehicle["driver_name"]
    ws["A6"] = "المشروع"
    ws["B6"] = vehicle["project"]
    ws["A7"] = "سعر النقلة"
    ws["B7"] = vehicle["trip_price"]

    ws["A9"] = "التاريخ"
    ws["B9"] = "المشروع"
    ws["C9"] = "من"
    ws["D9"] = "إلى"
    ws["E9"] = "نوع الحمولة"
    ws["F9"] = "عدد النقلات"
    ws["G9"] = "سعر النقلة"
    ws["H9"] = "إجمالي المستحق"
    ws["I9"] = "ملاحظات"
    style_header(ws, 9, 1, 9)
    rows = get_transport_records(vehicle["id"])
    for idx, row in enumerate(rows, start=10):
        ws[f"A{idx}"] = row["record_date"]
        ws[f"B{idx}"] = row["project"]
        ws[f"C{idx}"] = row["origin"]
        ws[f"D{idx}"] = row["destination"]
        ws[f"E{idx}"] = row["load_type"]
        ws[f"F{idx}"] = row["trips"]
        ws[f"G{idx}"] = row["trip_price"]
        ws[f"H{idx}"] = row["total_amount"]
        ws[f"I{idx}"] = row["notes"]
    ws.freeze_panes = "A10"


def export_water_detail(ws, truck):
    ws["A1"] = "الكود"
    ws["B1"] = truck["code"]
    ws["A2"] = "رقم العربية"
    ws["B2"] = truck["vehicle_number"]
    ws["A3"] = "رقم اللوحة"
    ws["B3"] = truck["plate_number"]
    ws["A4"] = "السائق"
    ws["B4"] = truck["driver_name"]
    ws["A5"] = "السعة"
    ws["B5"] = truck["capacity"]
    ws["A6"] = "المشروع"
    ws["B6"] = truck["project"]
    ws["A7"] = "سعر النقلة"
    ws["B7"] = truck["trip_price"]

    ws["A9"] = "التاريخ"
    ws["B9"] = "المشروع"
    ws["C9"] = "المكان"
    ws["D9"] = "نوع النقلة"
    ws["E9"] = "عدد النقلات"
    ws["F9"] = "سعر النقلة"
    ws["G9"] = "إجمالي المستحق"
    ws["H9"] = "ملاحظات"
    style_header(ws, 9, 1, 8)
    rows = get_water_records(truck["id"])
    for idx, row in enumerate(rows, start=10):
        ws[f"A{idx}"] = row["record_date"]
        ws[f"B{idx}"] = row["project"]
        ws[f"C{idx}"] = row["location"]
        ws[f"D{idx}"] = row["trip_type"]
        ws[f"E{idx}"] = row["trips"]
        ws[f"F{idx}"] = row["trip_price"]
        ws[f"G{idx}"] = row["total_amount"]
        ws[f"H{idx}"] = row["notes"]
    ws.freeze_panes = "A10"


def export_excel(file_path=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "الرئيسية"

    add_summary_sheet(wb)

    workers = list_employees()
    workers_ws = wb.create_sheet("العمال")
    workers_ws.append(["الكود", "اسم العامل", "الوظيفة", "الهاتف", "الرقم القومي", "المشروع", "سعر اليومية", "الحالة"])
    for emp in workers:
        workers_ws.append([
            emp["code"], emp["name"], emp["job_title"], emp["phone"], emp["national_id"],
            emp["project"], emp["daily_rate"], emp["status"]
        ])
    style_header(workers_ws, 1, 1, 8)
    workers_ws.freeze_panes = "A2"

    equipment_items = list_equipment()
    eq_ws = wb.create_sheet("المعدات")
    eq_ws.append(["الكود", "اسم المعدة", "النوع", "رقم المعدة", "المشغل", "المشروع", "سعر الساعة", "الحالة"])
    for item in equipment_items:
        eq_ws.append([
            item["code"], item["name"], item["equipment_type"], item["equipment_number"],
            item["operator_name"], item["project"], item["hourly_rate"], item["status"]
        ])
    style_header(eq_ws, 1, 1, 8)
    eq_ws.freeze_panes = "A2"

    transports = list_transport_vehicles()
    tr_ws = wb.create_sheet("عربيات النقل")
    tr_ws.append(["الكود", "رقم العربية", "نوع العربية", "اللوحة", "السائق", "المشروع", "سعر النقلة", "الحالة"])
    for item in transports:
        tr_ws.append([
            item["code"], item["vehicle_number"], item["vehicle_type"], item["plate_number"],
            item["driver_name"], item["project"], item["trip_price"], item["status"]
        ])
    style_header(tr_ws, 1, 1, 8)
    tr_ws.freeze_panes = "A2"

    waters = list_water_trucks()
    wt_ws = wb.create_sheet("عربيات المياه")
    wt_ws.append(["الكود", "رقم العربية", "اللوحة", "السائق", "السعة", "المشروع", "سعر النقلة", "الحالة"])
    for item in waters:
        wt_ws.append([
            item["code"], item["vehicle_number"], item["plate_number"], item["driver_name"],
            item["capacity"], item["project"], item["trip_price"], item["status"]
        ])
    style_header(wt_ws, 1, 1, 8)
    wt_ws.freeze_panes = "A2"

    for emp in workers:
        detail = wb.create_sheet(safe_sheet_name(f"{emp['code']} - {emp['name']}"))
        export_employee_detail(detail, emp)

    for item in equipment_items:
        detail = wb.create_sheet(safe_sheet_name(f"{item['code']} - {item['name']}"))
        export_equipment_detail(detail, item)

    for item in transports:
        detail = wb.create_sheet(safe_sheet_name(f"{item['code']} - {item['vehicle_number']}"))
        export_transport_detail(detail, item)

    for item in waters:
        detail = wb.create_sheet(safe_sheet_name(f"{item['code']} - {item['vehicle_number']}"))
        export_water_detail(detail, item)

    if file_path is None:
        from datetime import datetime
        file_path = OUTPUT_DIR / f"export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(file_path)
    return str(file_path)
